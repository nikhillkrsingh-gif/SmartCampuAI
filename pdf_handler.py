import os
import csv
import re
import numpy as np
from PIL import Image
from pypdf import PdfReader
from docx import Document

from config import Config
from database import get_db

ocr_reader = None

def get_ocr_reader():
    """Initializes and caches EasyOCR reader on CPU."""
    global ocr_reader
    if ocr_reader is None:
        try:
            import easyocr
            print("[OCR Info]: Initializing EasyOCR engine (English, CPU mode)...")
            ocr_reader = easyocr.Reader(['en'], gpu=False)
            print("[OCR Info]: EasyOCR engine loaded successfully.")
        except Exception as e:
            print(f"[OCR Warning]: Failed to initialize EasyOCR: {e}")
            ocr_reader = None
    return ocr_reader

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in Config.ALLOWED_EXTENSIONS

def clean_text(text):
    if not text:
        return ""
    text = str(text).replace('\x00', ' ')
    text = re.sub(r'[ \t\r\f]+', ' ', text)
    text = re.sub(r'\n\s*\n+', '\n\n', text)
    return text.strip()

def chunk_text_stream(full_text, chunk_size=200, overlap=35):
    words = full_text.split()
    if not words:
        return []
    if len(words) <= chunk_size:
        return [" ".join(words)]

    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i:i + chunk_size])
        if chunk.strip():
            chunks.append(chunk)
        i += (chunk_size - overlap)
    return chunks

def extract_and_chunk_image(filepath):
    """Extracts text from images using in-memory array conversion to bypass Windows filename bugs."""
    chunks = []
    reader = get_ocr_reader()
    if not reader:
        print("[OCR Error]: EasyOCR reader is unavailable. Ensure 'easyocr' is installed.")
        return []

    try:
        # Load image via PIL to safely bypass Windows path/spacing restrictions
        with Image.open(filepath) as pil_img:
            rgb_img = pil_img.convert("RGB")
            img_array = np.array(rgb_img)

        # Run text recognition on the raw pixel array
        results = reader.readtext(img_array, detail=0)
        extracted_text = clean_text(" ".join(results))

        if extracted_text:
            text_chunks = chunk_text_stream(extracted_text, chunk_size=200, overlap=35)
            for idx, c in enumerate(text_chunks):
                chunks.append((f"Image OCR Part {idx + 1:02d}", c))
        else:
            print(f"[OCR Warning]: No readable printed text found in '{os.path.basename(filepath)}'.")
    except Exception as e:
        print(f"[OCR Processing Error]: {e}")

    return chunks

def extract_and_chunk_pdf(filepath):
    chunks = []
    try:
        reader = PdfReader(filepath, strict=False)
        for idx, page in enumerate(reader.pages):
            page_num = idx + 1
            raw_text = page.extract_text() or ""
            cleaned = clean_text(raw_text)
            if cleaned:
                page_chunks = chunk_text_stream(cleaned, chunk_size=200, overlap=35)
                for c_idx, sub_c in enumerate(page_chunks):
                    chunks.append((f"Page {page_num:02d} Part {c_idx + 1:02d}", sub_c))
    except Exception as e:
        print(f"[PDF Extract Error]: {e}")
    return chunks

def extract_and_chunk_docx(filepath):
    chunks = []
    try:
        doc = Document(filepath)
        lines = []
        for p in doc.paragraphs:
            c = clean_text(p.text)
            if c:
                lines.append(c)
        for table in doc.tables:
            for row in table.rows:
                cells = [clean_text(cell.text) for cell in row.cells if clean_text(cell.text)]
                if cells:
                    lines.append(" | ".join(cells))
        combined = "\n".join(lines)
        if combined.strip():
            doc_chunks = chunk_text_stream(combined, chunk_size=200, overlap=35)
            for idx, c in enumerate(doc_chunks):
                chunks.append((f"Section {idx + 1:02d}", c))
    except Exception as e:
        print(f"[DOCX Extract Error]: {e}")
    return chunks

def extract_and_chunk_csv(filepath):
    chunks = []
    try:
        with open(filepath, mode='r', encoding='utf-8', errors='ignore') as f:
            reader = list(csv.reader(f))
            if not reader:
                return []
            header = " | ".join([clean_text(c) for c in reader[0] if clean_text(c)])
            data_rows = reader[1:]
            batch_size = 12
            for i in range(0, len(data_rows), batch_size):
                batch = data_rows[i:i + batch_size]
                rows_str = [" | ".join([clean_text(c) for c in r if clean_text(c)]) for r in batch]
                formatted = [r for r in rows_str if r]
                if formatted:
                    chunk_text = f"Headers: {header}\n" + "\n".join(formatted)
                    chunks.append((f"Rows {i + 1:03d}-{i + len(formatted):03d}", chunk_text))
    except Exception as e:
        print(f"[CSV Extract Error]: {e}")
    return chunks

def index_campus_file(filename, uploader_role):
    """Routes the uploaded file to its format parser, creates chunks, and stores them in SQLite."""
    filepath = os.path.join(Config.UPLOAD_FOLDER, filename)
    ext = filename.rsplit('.', 1)[-1].lower()

    if ext == 'pdf':
        chunks = extract_and_chunk_pdf(filepath)
    elif ext == 'docx':
        chunks = extract_and_chunk_docx(filepath)
    elif ext == 'csv':
        chunks = extract_and_chunk_csv(filepath)
    elif ext in ['png', 'jpg', 'jpeg']:
        chunks = extract_and_chunk_image(filepath)
    else:
        return False, "Unsupported format."

    if not chunks:
        return False, "No readable text could be extracted from this file."

    conn = get_db()
    c = conn.cursor()
    c.execute("DELETE FROM document_vectors WHERE filename = ?", (filename,))

    for sec_label, chunk_text in chunks:
        c.execute('''INSERT INTO document_vectors (filename, section_label, chunk_text, uploader_role, access_level)
                     VALUES (?, ?, ?, ?, 'public_student')''',
                  (filename, sec_label, chunk_text, uploader_role))

    c.execute('''INSERT INTO campus_documents (filename, total_chunks, uploader_role, access_level)
                 VALUES (?, ?, ?, 'public_student')
                 ON CONFLICT(filename) DO UPDATE SET 
                 total_chunks = excluded.total_chunks,
                 uploader_role = excluded.uploader_role,
                 access_level = 'public_student',
                 uploaded_at = CURRENT_TIMESTAMP''',
              (filename, len(chunks), uploader_role))

    conn.commit()
    conn.close()
    return True, f"Indexed '{filename}' ({len(chunks)} chunks created)."

def retrieve_best_chunks(query, user_role, top_k=5):
    """Searches SQLite for relevant document chunks using composite scoring across filenames, sections, and text."""
    conn = get_db()
    c = conn.cursor()
    rows = c.execute("SELECT filename, section_label, chunk_text FROM document_vectors").fetchall()
    conn.close()

    if not rows:
        return "NO_OFFICIAL_DATA"

    query_lower = query.lower()
    stop_words = {
        "what", "tell", "give", "show", "from", "when", "where", "how", "is", 
        "are", "the", "in", "of", "about", "for", "any", "can", "you", "please", 
        "and", "to", "a", "an", "do", "does", "me", "my", "our", "all"
    }
    raw_tokens = re.findall(r'\b\w{2,}\b', query_lower)
    tokens = [w for w in raw_tokens if w not in stop_words] or raw_tokens

    if not tokens:
        return "NO_OFFICIAL_DATA"

    # Check if query specifically asks for a roster/list/all entries
    is_roster_query = any(w in query_lower for w in [
        "list", "roster", "all team", "teams", "shortlist", "participants", 
        "candidates", "who are selected", "roll number"
    ])

    scored = []
    for r in rows:
        fn_lower = r["filename"].lower()
        sec_lower = r["section_label"].lower()
        txt_lower = r["chunk_text"].lower()

        score = 0
        for token in tokens:
            if token in fn_lower:
                score += 25
            if token in sec_lower:
                score += 15
            cnt = txt_lower.count(token)
            if cnt > 0:
                score += min(cnt, 8) * 5

        # Phrase / bigram match bonus
        for idx in range(len(tokens) - 1):
            bigram = f"{tokens[idx]} {tokens[idx + 1]}"
            if bigram in fn_lower:
                score += 50
            if bigram in txt_lower:
                score += 30

        if score > 0:
            scored.append((score, r["filename"], r["section_label"], r["chunk_text"]))

    if not scored:
        return "NO_OFFICIAL_DATA"

    scored.sort(key=lambda x: x[0], reverse=True)

    # For roster queries, ensure all sequential chunks of the top-matching document are included
    if is_roster_query and scored:
        top_doc = scored[0][1]
        doc_chunks = [r for r in rows if r["filename"] == top_doc]
        if len(doc_chunks) <= 8:
            return "\n\n".join([f"=== Document Record: {r['filename']} ({r['section_label']}) ===\n{r['chunk_text']}" for r in doc_chunks])

    top_matches = scored[:top_k]
    return "\n\n".join([f"=== Document Record: {fn} ({sec}) ===\n{ct}" for _, fn, sec, ct in top_matches])